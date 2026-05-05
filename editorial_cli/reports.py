from __future__ import annotations

import datetime as dt
import html as html_module
import json
import re
from pathlib import Path

from editorial_cli.models import JsonValue, Section, Suggestion


def render_markdown_report(source_name: str, context_brief: str, suggestions: list[Suggestion]) -> str:
    lines = [
        f"# Editorial Suggestions for {source_name}",
        "",
        f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
        "## Whole-Manuscript Context and Style Brief",
        "",
    ]
    lines.extend(f"> {line}" if line else ">" for line in context_brief.splitlines())
    lines.append("")

    for item in suggestions:
        lines.extend(["", f"## {item.get('title', 'Untitled Section')}", ""])
        summary = str(item.get("summary") or "").strip()
        if summary:
            lines.extend(["**Summary:** " + summary, ""])
        add_markdown_list(lines, "Suggestions", item.get("suggestions"))
        add_markdown_list(lines, "Style Preservation", item.get("style_preservation"))
        add_markdown_list(lines, "Continuity", item.get("continuity"))
        add_markdown_list(lines, "Line-Level Notes", item.get("line_level"))

    return "\n".join(lines).strip() + "\n"


def render_json_report(source_name: str, context_brief: str, suggestions: list[Suggestion]) -> str:
    payload = {
        "source": source_name,
        "generated_at": dt.datetime.now().replace(microsecond=0).isoformat(),
        "context_brief": context_brief,
        "sections": suggestions,
    }
    return json.dumps(payload, indent=2) + "\n"


def render_markdown_section_report(source_name: str, suggestion: Suggestion) -> str:
    title = str(suggestion.get("title") or "Untitled Section").strip() or "Untitled Section"
    lines = [
        f"# {title}",
        "",
        f"Source: {source_name}",
        "",
    ]
    summary = str(suggestion.get("summary") or "").strip()
    if summary:
        lines.extend(["**Summary:** " + summary, ""])
    add_markdown_list(lines, "Suggestions", suggestion.get("suggestions"))
    add_markdown_list(lines, "Style Preservation", suggestion.get("style_preservation"))
    add_markdown_list(lines, "Continuity", suggestion.get("continuity"))
    add_markdown_list(lines, "Line-Level Notes", suggestion.get("line_level"))
    return "\n".join(lines).strip() + "\n"


def render_markdown_index(source_name: str, suggestions: list[Suggestion], filenames: list[str]) -> str:
    lines = [
        f"# Editorial Suggestions for {source_name}",
        "",
        f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
        "## Files",
        "",
        "- [Whole-Manuscript Context and Style Brief](context_brief.md)",
    ]
    for suggestion, filename in zip(suggestions, filenames, strict=False):
        title = str(suggestion.get("title") or "Untitled Section").strip() or "Untitled Section"
        lines.append(f"- [{title}]({filename})")
    return "\n".join(lines).strip() + "\n"


def write_markdown_report_directory(
    output_dir: Path,
    source_name: str,
    context_brief: str,
    suggestions: list[Suggestion],
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    filenames = section_report_filenames(suggestions)
    (output_dir / "context_brief.md").write_text(render_context_brief(context_brief), encoding="utf-8")
    (output_dir / "index.md").write_text(render_markdown_index(source_name, suggestions, filenames), encoding="utf-8")
    for suggestion, filename in zip(suggestions, filenames, strict=False):
        (output_dir / filename).write_text(render_markdown_section_report(source_name, suggestion), encoding="utf-8")
    return output_dir


def render_context_brief(context_brief: str) -> str:
    lines = ["# Whole-Manuscript Context and Style Brief", ""]
    lines.extend(context_brief.splitlines())
    return "\n".join(lines).strip() + "\n"


def section_report_filenames(suggestions: list[Suggestion]) -> list[str]:
    seen: dict[str, int] = {}
    filenames: list[str] = []
    for index, suggestion in enumerate(suggestions, start=1):
        title = str(suggestion.get("title") or "untitled-section")
        slug = slugify_filename(title)
        seen[slug] = seen.get(slug, 0) + 1
        suffix = f"-{seen[slug]}" if seen[slug] > 1 else ""
        filenames.append(f"{index:03d}-{slug}{suffix}.md")
    return filenames


def render_report(
    source_name: str,
    context_brief: str,
    suggestions: list[Suggestion],
    output_format: str,
) -> str:
    if output_format == "json":
        return render_json_report(source_name, context_brief, suggestions)
    if output_format == "html":
        return render_html_report(source_name, context_brief, suggestions)
    return render_markdown_report(source_name, context_brief, suggestions)


def render_outline(source_name: str, sections: list[Section], output_format: str = "text") -> str:
    if output_format == "json":
        payload = {
            "source": source_name,
            "section_count": len(sections),
            "sections": [
                {
                    "index": section.index,
                    "title": section.title,
                    "characters": len(section.text),
                    "preview": preview_text(section.text),
                }
                for section in sections
            ],
        }
        return json.dumps(payload, indent=2) + "\n"

    if output_format == "markdown":
        lines = ["# Manuscript Outline", "", f"Source: {source_name}", "", f"Sections: {len(sections)}", ""]
        for section in sections:
            lines.extend(
                [
                    f"## {section.index}. {section.title}",
                    "",
                    f"- Characters: {len(section.text)}",
                    f"- Preview: {preview_text(section.text)}",
                    "",
                ]
            )
        return "\n".join(lines).strip() + "\n"

    width = max([len(section.title) for section in sections] + [5])
    lines = [f"Manuscript outline for {source_name}", f"Sections: {len(sections)}", ""]
    for section in sections:
        lines.append(
            f"{section.index:>3}. {section.title:<{width}}  "
            f"{len(section.text):>7} chars  {preview_text(section.text)}"
        )
    return "\n".join(lines).strip() + "\n"


def preview_text(text: str, limit: int = 120) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    if len(compact) <= limit:
        return compact
    return compact[: limit - 3].rstrip() + "..."


def slugify_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value.strip().lower()).strip("-")
    return cleaned or "untitled-section"


def render_html_report(source_name: str, context_brief: str, suggestions: list[Suggestion]) -> str:
    generated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    e = html_module.escape

    def html_list(title: str, values: JsonValue) -> str:
        if not values:
            return ""
        if isinstance(values, str):
            values = [values]
        if not isinstance(values, list):
            return ""
        items = [str(v).strip() for v in values if str(v).strip()]
        if not items:
            return ""
        lis = "".join(f"<li>{e(item)}</li>" for item in items)
        return f"<h4>{e(title)}</h4><ul>{lis}</ul>"

    sections_html = ""
    for item in suggestions:
        title = str(item.get("title") or "Untitled Section")
        summary = str(item.get("summary") or "").strip()
        summary_html = f"<p class='summary'><strong>Summary:</strong> {e(summary)}</p>" if summary else ""
        body = (
            summary_html
            + html_list("Suggestions", item.get("suggestions"))
            + html_list("Style Preservation", item.get("style_preservation"))
            + html_list("Continuity", item.get("continuity"))
            + html_list("Line-Level Notes", item.get("line_level"))
        )
        sections_html += f"""
    <details open>
      <summary class='section-title'>{e(title)}</summary>
      <div class='section-body'>{body}</div>
    </details>"""

    brief_paragraphs = "".join(f"<p>{e(line)}</p>" if line.strip() else "" for line in context_brief.splitlines())

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Editorial Suggestions — {e(source_name)}</title>
<style>
  :root {{ --bg:#fff; --fg:#1a1a1a; --muted:#555; --accent:#0057b8; --border:#ddd; --code-bg:#f5f5f5; }}
  @media (prefers-color-scheme:dark) {{
    :root {{ --bg:#121212; --fg:#e8e8e8; --muted:#aaa; --accent:#5ba3f5; --border:#333; --code-bg:#1e1e1e; }}
  }}
  *,*::before,*::after {{ box-sizing:border-box; }}
  body {{ font-family:Georgia,serif; background:var(--bg); color:var(--fg); max-width:860px; margin:0 auto; padding:2rem 1.5rem; line-height:1.7; }}
  h1 {{ font-size:1.6rem; margin-bottom:.25rem; }}
  .meta {{ color:var(--muted); font-size:.85rem; margin-bottom:2rem; }}
  details {{ border:1px solid var(--border); border-radius:6px; margin-bottom:1rem; }}
  details[open] {{ background:var(--code-bg); }}
  summary {{ cursor:pointer; padding:.75rem 1rem; font-weight:bold; list-style:none; }}
  summary::-webkit-details-marker {{ display:none; }}
  summary.section-title {{ font-size:1.05rem; color:var(--accent); }}
  .brief-body {{ padding:.75rem 1rem; font-size:.9rem; color:var(--muted); }}
  .brief-body p {{ margin:.4rem 0; }}
  .section-body {{ padding:.5rem 1.25rem 1rem; }}
  .summary {{ margin-bottom:.75rem; }}
  h4 {{ font-size:.9rem; text-transform:uppercase; letter-spacing:.05em; color:var(--muted); margin:1rem 0 .4rem; }}
  ul {{ margin:0; padding-left:1.4rem; }}
  li {{ margin:.3rem 0; }}
</style>
</head>
<body>
<h1>Editorial Suggestions</h1>
<p class="meta">Source: {e(source_name)} &nbsp;·&nbsp; Generated: {e(generated)}</p>
<details>
  <summary>Whole-Manuscript Context and Style Brief</summary>
  <div class="brief-body">{brief_paragraphs}</div>
</details>
{sections_html}
</body>
</html>
"""


def render_compare_report(
    id1: str,
    id2: str,
    suggestions1: list[Suggestion],
    suggestions2: list[Suggestion],
    output_format: str = "text",
) -> str:
    index1 = {str(s.get("title", "")): s for s in suggestions1}
    index2 = {str(s.get("title", "")): s for s in suggestions2}
    all_titles = list(dict.fromkeys(list(index1) + list(index2)))

    if output_format == "json":
        rows = []
        for title in all_titles:
            s1, s2 = index1.get(title), index2.get(title)
            rows.append({
                "title": title,
                "status": "added" if s1 is None else "removed" if s2 is None else "changed" if s1 != s2 else "unchanged",
                "run1": s1,
                "run2": s2,
            })
        return json.dumps({"run1": id1, "run2": id2, "sections": rows}, indent=2) + "\n"

    lines = [f"Comparing {id1} → {id2}", ""]
    for title in all_titles:
        s1, s2 = index1.get(title), index2.get(title)
        if s1 is None:
            lines += [f"+ {title}  [new in {id2}]", ""]
            continue
        if s2 is None:
            lines += [f"- {title}  [removed in {id2}]", ""]
            continue
        if s1 == s2:
            lines += [f"  {title}  [unchanged]", ""]
            continue
        lines += [f"~ {title}  [changed]"]
        sug1 = [str(x) for x in (s1.get("suggestions") or []) if x]
        sug2 = [str(x) for x in (s2.get("suggestions") or []) if x]
        removed = [x for x in sug1 if x not in sug2]
        added = [x for x in sug2 if x not in sug1]
        for note in removed:
            lines.append(f"  - {note}")
        for note in added:
            lines.append(f"  + {note}")
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def render_diff_report(
    name1: str,
    name2: str,
    sections1: list[Section],
    sections2: list[Section],
    output_format: str = "text",
) -> str:
    import hashlib

    def _hash(s: Section) -> str:
        return hashlib.md5(s.text.encode()).hexdigest()

    index1 = {s.title: s for s in sections1}
    index2 = {s.title: s for s in sections2}
    all_titles = list(dict.fromkeys(list(index1) + list(index2)))

    rows = []
    for title in all_titles:
        s1, s2 = index1.get(title), index2.get(title)
        if s1 is None:
            status = "added"
        elif s2 is None:
            status = "removed"
        elif _hash(s1) != _hash(s2):
            status = "changed"
        else:
            status = "unchanged"
        rows.append({"title": title, "status": status, "chars1": len(s1.text) if s1 else 0, "chars2": len(s2.text) if s2 else 0})

    if output_format == "json":
        return json.dumps({"file1": name1, "file2": name2, "sections": rows}, indent=2) + "\n"

    lines = [f"Structural diff: {name1} → {name2}", ""]
    counts = {"added": 0, "removed": 0, "changed": 0, "unchanged": 0}
    for row in rows:
        status = str(row["status"])
        counts[status] = counts.get(status, 0) + 1
        prefix = {"added": "+", "removed": "-", "changed": "~", "unchanged": " "}.get(status, " ")
        chars = f"  ({row['chars2']} chars)" if status in {"added", "unchanged"} else f"  ({row['chars1']} → {row['chars2']} chars)"
        lines.append(f"{prefix} {row['title']}{chars}")
    lines += [
        "",
        f"{counts['added']} added  {counts['removed']} removed  "
        f"{counts['changed']} changed  {counts['unchanged']} unchanged",
    ]
    return "\n".join(lines).strip() + "\n"


def write_docx_report(
    output: Path,
    source_name: str,
    context_brief: str,
    suggestions: list[Suggestion],
) -> None:
    try:
        from docx import Document  # type: ignore[import-untyped]
        from docx.shared import Pt, RGBColor  # type: ignore[import-untyped]
    except ImportError as exc:
        from editorial_cli.errors import CliError
        raise CliError(
            "python-docx is required for .docx output.\n"
            "Install it with:  pip install python-docx"
        ) from exc

    doc = Document()
    doc.add_heading(f"Editorial Suggestions for {source_name}", 0)
    meta = doc.add_paragraph(f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    meta.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_heading("Whole-Manuscript Context and Style Brief", 1)
    for line in context_brief.splitlines():
        if line.strip():
            p = doc.add_paragraph(style="Quote")
            p.add_run(line)

    for item in suggestions:
        title = str(item.get("title") or "Untitled Section")
        doc.add_heading(title, 1)
        summary = str(item.get("summary") or "").strip()
        if summary:
            p = doc.add_paragraph()
            p.add_run("Summary: ").bold = True
            p.add_run(summary)

        def _add_list(heading: str, values: JsonValue) -> None:
            if not values:
                return
            if isinstance(values, str):
                values = [values]
            if not isinstance(values, list):
                return
            items = [str(v).strip() for v in values if str(v).strip()]
            if not items:
                return
            p2 = doc.add_paragraph()
            p2.add_run(f"{heading}:").bold = True
            p2.runs[0].font.size = Pt(10)
            for note in items:
                doc.add_paragraph(note, style="List Bullet")

        _add_list("Suggestions", item.get("suggestions"))
        _add_list("Style Preservation", item.get("style_preservation"))
        _add_list("Continuity", item.get("continuity"))
        _add_list("Line-Level Notes", item.get("line_level"))

    doc.save(output)


def add_markdown_list(lines: list[str], title: str, values: JsonValue) -> None:
    if not values:
        return
    if isinstance(values, str):
        values = [values]
    if not isinstance(values, list):
        return
    cleaned = [str(value).strip() for value in values if str(value).strip()]
    if not cleaned:
        return
    lines.extend([f"**{title}:**", ""])
    lines.extend(f"- {value}" for value in cleaned)
    lines.append("")
